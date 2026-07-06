#!/usr/bin/env python3
"""Import monthly cleaning-schedule spreadsheets into data/jobs.seed.json.

Usage:
    python scripts/import_schedule.py schedule_MAY.xlsx schedule_JUNE.xlsx [...]

Each spreadsheet is the agency's "Cleaning Schedule & Assignments" export with a
header row at row 3 and columns:
    A Name(code+customer)  C Contact  D Date  E Time  F Hours  G Assign Person
    H House/Location  I Location(postal)  J Job Status  K Payment Status  L Remark

Requires: openpyxl  (pip install openpyxl)
"""
import sys, re, json, os
import openpyxl

RATE = 17  # default S$/hour — keep in sync with COMPANY.hourlyRate

CLEANER_MAP = {
    "nant thin zar": "cln-ntz",
    "yu ya naing": "cln-yyn",
    "hnin hnin": "cln-hh",
    "sai aye": "cln-sa",
    "thidar win": "cln-tw",
    "jenny": "cln-jen",
}


def parse_hours(f):
    if f is None:
        return None
    m = re.search(r"(\d+(?:\.\d+)?)", str(f))
    return float(m.group(1)) if m else None


def parse_start(e):
    if e is None:
        return None
    part = str(e).upper().replace(" ", "").split("-")[0]
    m = re.match(r"(\d{1,2})(?:[.:](\d{2}))?(AM|PM)?", part)
    if not m:
        return None
    hh, mm, ap = int(m.group(1)), int(m.group(2) or 0), m.group(3)
    if ap == "PM" and hh != 12:
        hh += 12
    if ap == "AM" and hh == 12:
        hh = 0
    if not ap and hh < 7:
        hh += 12
    return min(hh, 23) * 60 + mm


def to_hhmm(mins):
    mins = max(0, min(mins, 23 * 60 + 59))
    return f"{mins // 60:02d}:{mins % 60:02d}"


def map_status(j):
    v = (str(j) if j else "").strip().lower()
    return "completed" if v.startswith("complet") else "cancelled" if v.startswith("cancel") else "scheduled"


def map_payment(k):
    v = (str(k) if k else "").strip().lower()
    return "done" if "done" in v else "pending" if "pending" in v else "unbilled"


def split_customer(a):
    a = str(a).strip()
    m = re.match(r"^\s*((?:MCC|MC)\s*\d+)\s*[-: ]*\s*(.*)$", a, re.I)
    if m:
        return re.sub(r"\s+", "", m.group(1)).upper(), (m.group(2).strip() or re.sub(r"\s+", "", m.group(1)).upper())
    return a, a


def clean_addr(h):
    return re.sub(r"\s*\n\s*", ", ", str(h).strip()) if h else ""


def main(paths):
    customers, jobs = {}, []
    for path in paths:
        tag = re.sub(r"[^a-z]", "", os.path.basename(path).lower())[:6]
        ws = openpyxl.load_workbook(path, data_only=True).active
        for r in range(4, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if not a:
                continue
            cleaner = ws.cell(r, 7).value
            cid = CLEANER_MAP.get(str(cleaner).strip().lower()) if cleaner else None
            date = ws.cell(r, 4).value
            if cid is None or date is None:
                continue
            hours = parse_hours(ws.cell(r, 6).value) or 3.0
            start = parse_start(ws.cell(r, 5).value)
            start = 9 * 60 if start is None else start
            code, name = split_customer(a)
            c = customers.setdefault(code, {"id": "cus-" + code.lower(), "code": code, "name": name,
                                            "phone": "", "email": "", "address": "", "postal": ""})
            if len(name) > len(c["name"]):
                c["name"] = name
            if ws.cell(r, 3).value:
                c["phone"] = str(ws.cell(r, 3).value).strip()
            if ws.cell(r, 8).value:
                c["address"] = clean_addr(ws.cell(r, 8).value)
            if ws.cell(r, 9).value:
                c["postal"] = str(ws.cell(r, 9).value).strip()
            status = map_status(ws.cell(r, 10).value)
            jobs.append({
                "id": f"job-{tag}-{r}",
                "customerId": c["id"],
                "cleanerId": cid,
                "date": date.strftime("%Y-%m-%d"),
                "startTime": to_hhmm(start),
                "endTime": to_hhmm(start + int(round(hours * 60))),
                "hours": round(hours, 2),
                "amount": int(round(hours * RATE)),
                "jobStatus": status,
                "paymentStatus": map_payment(ws.cell(r, 11).value),
                "remark": (str(ws.cell(r, 12).value).strip() if ws.cell(r, 12).value else ""),
                "emailSent": status == "completed",
            })

    out_path = os.path.join(os.path.dirname(__file__), "..", "data", "jobs.seed.json")
    with open(out_path, "w") as fh:
        json.dump({"rate": RATE, "customers": list(customers.values()), "jobs": jobs}, fh, indent=1, ensure_ascii=False)
    print(f"Wrote {len(jobs)} jobs / {len(customers)} customers to data/jobs.seed.json")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Usage: python scripts/import_schedule.py <xlsx> [<xlsx> ...]")
    main(sys.argv[1:])
